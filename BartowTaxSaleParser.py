from xlwt import Workbook
import docx

class BartowTaxSaleParser:

    #source: file name string (must be word doc) containing the tax sale information
    #file: string that represents the keyword to parse for the file number 
    #parcel: string that represents the keyword to parse for the parcel number i.e. Map/Parcel Number:
    #owner: string that represents the keyword to parse for the owner i.e. "Current Property Owner"
    #address: string that represents the keyword to parse for the address i.e. "known as"
    #output: output file name
    def __init__(self, source: str, file: str, parcel: str, owner: str, address: str,output: str, second_address="located on") -> None:
        self.text = self.readtxt(source)
        self.file = file
        self.parcel = parcel
        self.owner = owner
        self.address = address
        self.second_address = second_address
        self.output_file_name = output
        self.info_list = self.split_by_file()
        self.file_number_list = []
        self.parcel_list = []
        self.owner_list = []
        self.address_list = []

    def retrieve_file(self):
        for info in self.info_list:
            file_number = self.find_text(info, ": ", "Map", 0)
            self.file_number_list.append(file_number)

        print(self.file_number_list)


    def retrieve_parcel(self):
        for info in self.info_list:
            self.parcel_list.append(self.find_text(info, self.parcel, "Defendant", 2))
        print(self.parcel_list)
        print(len(self.parcel_list))

    def retrieve_owner(self, end="Reference Deed",fifa_reference="Defendant(s) in FiFa"):
        for info in self.info_list:
            owner = self.find_text(info, self.owner, end, 2)
            if (owner == "Same as Defendant(s) in FiFa"):
                owner = self.find_text(info, fifa_reference, ";", 2)
            self.owner_list.append(owner)
        print(self.owner_list)
        print(len(self.owner_list)) 
    
    def retrieve_address(self):
        for info in self.info_list:
            address = self.find_text(info, self.address, ".", 1)
            if (address is None):
                address = self.find_text(info, self.second_address, ".", 1)
                if (address is None): # this will get most of them
                    self.address_list.append("Address or street not found")
                    continue
            self.address_list.append(address)
        print(self.address_list)
        print(len(self.address_list)) 

    def revise_info(self):
        for info in self.info_list:
            info = "File #" + info
    def generate_excel(self):
        workbook = Workbook()
        sheet1 = workbook.add_sheet("Tax Sales")
        sheet1.write(0, 0, 'File #')
        sheet1.write(0, 1, 'Map/Parcel Number')
        sheet1.write(0, 2, 'Owner')
        sheet1.write(0, 3, 'Address or Location')
        sheet1.write(0, 4, 'Additional info')
        row_offset = 1
        self.revise_info()
        all_lists = (self.file_number_list, self.parcel_list, self.owner_list, self.address_list, self.info_list)
        for col, list in enumerate(all_lists):
            for row, val in enumerate(list):
                sheet1.write(row + row_offset, col, val)
        workbook.save(f"../Result/{self.output_file_name}.xls")
        #
            


    def find_text(self, text, start_text, end_text, offset):
        index_of_address = text.find(start_text)
        if index_of_address == -1:
            return None
        start_index = index_of_address + len(start_text) + offset
        end_index = start_index + text[start_index::].find(end_text)
        return text[start_index:end_index].strip()

    def split_by_file(self):
        info_list = self.text.split(self.file)
        info_list.pop(0)
        return info_list

    def readtxt(self, filename):
        doc = docx.Document(filename)
        fullText = []
        for para in doc.paragraphs:
            fullText.append(para.text)
        return ' '.join(fullText)

parser = BartowTaxSaleParser("../Source/source.docx", "File #", "Map/Parcel Number", "Current Property Owner", "known as", "result")
# parser.retrieve_parcel()
parser.retrieve_file()
parser.retrieve_parcel()
parser.retrieve_owner()
parser.retrieve_address()
parser.generate_excel()